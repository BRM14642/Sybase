from docx import Document
import os
import subprocess
from src.utils.logging_config import logger

class DocumentHandler:
    def __init__(self):
        self.file_path = ''

    def fill_template(self, template_path, output_path, replacements):
        # Abre el documento de la plantilla
        doc = Document(template_path)

        # Reemplaza el texto en las secciones especificadas manteniendo el formato
        for paragraph in doc.paragraphs:
            for key, value in replacements.items():
                if key in paragraph.text:
                    for run in paragraph.runs:
                        if key in run.text:
                            run.text = run.text.replace(key, value)

        # Reemplaza el texto en las tablas especificadas manteniendo el formato
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for key, value in replacements.items():
                            if key in paragraph.text:
                                for run in paragraph.runs:
                                    if key in run.text:
                                        run.text = run.text.replace(key, value)

        # Guarda el documento con los cambios
        doc.save(output_path)

    def convert_docx_to_pdf(self, docx_path, output_path):
        try:
            if not os.path.exists(docx_path):
                logger.info(f"Error: The file {docx_path} does not exist.")
                return


            logger.info(f"Converting {docx_path} to {output_path}...")
            subprocess.run([
                "/Applications/LibreOffice.app/Contents/MacOS/soffice",
                "--headless",
                "--convert-to", "pdf",
                docx_path,
                "--outdir", output_path
            ])

            logger.info(f"File converted to PDF: {output_path}")
        except Exception as e:
            logger.info(f"An error occurred during PDF conversion: {e}")